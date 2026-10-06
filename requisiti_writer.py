"""
requisiti_writer.py
-------------------
Compila il foglio "Requisiti" replicando la struttura gerarchica
del foglio "Funzioni".

  LIVELLO 1 — riga funzione (già esistente nel foglio)
      Col A: FUNC ID       "LV_Country_code_selection"
      Col B: DESCRIZIONE   "Selezione su IDU del paese corrente..."

  LIVELLO 2 — righe requisito (INSERITE dal codice sotto la funzione)
      Col A: vuota   ← discriminatore
      Col B: vuota
      Col C: Nr requisito
      Col D: Configurazione di appartenenza, colorata

FUNZIONAMENTO
  1. Cerca la riga della funzione scorrendo la colonna A (FUNC ID).
     Fallback: confronto sulla DESCRIZIONE (col B) se il FUNC ID non è
     trovato, utile quando l'ID è scritto in modo leggermente diverso.
  2. Determina la fine del blocco di secondo livello (prima riga con
     col A valorizzata, oppure fine foglio).
  3. I requisiti già presenti nel blocco vengono AGGIORNATI sul posto
     (idempotenza); quelli nuovi vengono INSERITI in coda al blocco
     con ws.insert_rows(), così le funzioni successive scalano in basso.
  4. Dopo ogni inserimento l'indice interno viene riallineato: senza
     questo passaggio le funzioni successive verrebbero scritte nelle
     righe sbagliate.
"""
from __future__ import annotations
import logging

from openpyxl.styles import Font, Alignment

import config as cfg
import config_colors

log = logging.getLogger(__name__)

try:
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    _RICH = True
except Exception:          # openpyxl < 3.1 → nessun rich text
    _RICH = False


class RequisitiWriter:
    """Scrive i requisiti come righe di secondo livello nel foglio Requisiti."""

    def __init__(self, ws):
        self.ws = ws
        # {func_id_normalizzato: riga di livello 1}
        self.func_rows: dict[str, int] = {}
        # {descrizione_normalizzata: riga di livello 1}
        self.desc_rows: dict[str, int] = {}
        self.rows_written = 0
        self._scan_level1()
        log.info(
            f"  [Requisiti] Foglio '{self.ws.title}': "
            f"{len(self.func_rows)} funzioni di livello 1 trovate"
        )

    # ─────────────────────────────────────────────────────────────────────
    # Indicizzazione livello 1
    # ─────────────────────────────────────────────────────────────────────

    def _scan_level1(self) -> None:
        self.func_rows.clear()
        self.desc_rows.clear()
        for row in range(cfg.REQ_HEADER_ROW + 1, self.ws.max_row + 1):
            fid = _s(self.ws.cell(row, cfg.REQ_FUNC_ID_COL).value)
            if not fid:
                continue
            self.func_rows.setdefault(_norm(fid), row)
            desc = _s(self.ws.cell(row, cfg.REQ_DESC_COL).value)
            if desc:
                self.desc_rows.setdefault(_norm(desc), row)

    def _shift(self, from_row: int, amount: int) -> None:
        """Riallinea l'indice dopo un inserimento di righe."""
        if amount <= 0:
            return
        for d in (self.func_rows, self.desc_rows):
            for k, v in d.items():
                if v >= from_row:
                    d[k] = v + amount

    # ─────────────────────────────────────────────────────────────────────
    # Ricerca / creazione riga funzione
    # ─────────────────────────────────────────────────────────────────────

    def _find_function_row(self, func_id: str, func_desc: str) -> int | None:
        row = self.func_rows.get(_norm(func_id))
        if row:
            return row
        row = self.desc_rows.get(_norm(func_desc))
        if row:
            log.info(
                f"  [Requisiti] '{func_id}' non trovato per FUNC ID — "
                f"abbinato per DESCRIZIONE (riga {row})"
            )
            return row
        return None

    def _create_function_row(self, func_id: str, func_desc: str) -> int:
        row = self.ws.max_row + 1
        c = self.ws.cell(row, cfg.REQ_FUNC_ID_COL)
        c.value = func_id
        c.font = Font(bold=True)
        d = self.ws.cell(row, cfg.REQ_DESC_COL)
        d.value = func_desc
        d.alignment = Alignment(wrap_text=True, vertical="top")
        self.func_rows[_norm(func_id)] = row
        if func_desc:
            self.desc_rows.setdefault(_norm(func_desc), row)
        log.info(f"  [Requisiti] Creata riga funzione '{func_id}' (riga {row})")
        return row

    # ─────────────────────────────────────────────────────────────────────
    # Blocco di secondo livello
    # ─────────────────────────────────────────────────────────────────────

    def _block_end(self, func_row: int) -> int:
        """
        Ultima riga del blocco di secondo livello della funzione.
        Il blocco termina alla prima riga con col A valorizzata
        (= funzione successiva) oppure a fine foglio.
        """
        row = func_row + 1
        last = self.ws.max_row
        while row <= last:
            if _s(self.ws.cell(row, cfg.REQ_FUNC_ID_COL).value):
                return row - 1
            row += 1
        return last

    def _existing_reqs(self, func_row: int, block_end: int) -> dict[str, int]:
        """{nr_requisito_normalizzato: riga} già presenti nel blocco."""
        out = {}
        for row in range(func_row + 1, block_end + 1):
            nr = _s(self.ws.cell(row, cfg.REQ_NUMBER_COL).value)
            if nr:
                out[_norm(_strip_progressive(nr))] = row
        return out

    # ─────────────────────────────────────────────────────────────────────
    # API pubblica
    # ─────────────────────────────────────────────────────────────────────

    def write_group(
        self,
        func_id: str,
        func_desc: str,
        reqs_by_config: dict,
    ) -> int:
        """
        reqs_by_config: {config_name: [Requirement, ...]}

        Scrive/aggiorna le righe di secondo livello della funzione.
        Ritorna il numero di righe scritte o aggiornate.
        """
        entries = self._merge(reqs_by_config)
        if not entries:
            return 0

        func_row = self._find_function_row(func_id, func_desc)
        if func_row is None:
            if not cfg.REQ_CREATE_MISSING_FUNCTION:
                log.warning(
                    f"  [Requisiti] Funzione '{func_id}' non presente nel foglio — "
                    "requisiti non scritti"
                )
                return 0
            func_row = self._create_function_row(func_id, func_desc)

        block_end = self._block_end(func_row)
        existing  = self._existing_reqs(func_row, block_end)

        aggiornati, nuovi = [], []
        for e in entries:
            if _norm(e["req_id"]) in existing:
                aggiornati.append(e)
            else:
                nuovi.append(e)

        n = 0

        # ── Aggiorna le righe già presenti ────────────────────────────────
        for e in aggiornati:
            row = existing[_norm(e["req_id"])]
            self._fill_row(row, e, progressive=None)
            n += 1

        # ── Inserisce le righe nuove in coda al blocco ────────────────────
        if nuovi:
            insert_at = block_end + 1
            self.ws.insert_rows(insert_at, amount=len(nuovi))
            # Le funzioni sottostanti sono scese: riallinea l'indice
            self._shift(from_row=insert_at, amount=len(nuovi))

            base = len(existing)
            for i, e in enumerate(nuovi):
                row = insert_at + i
                prog = base + i + 1 if cfg.REQ_PROGRESSIVE_NUMBER else None
                self._fill_row(row, e, progressive=prog)
                n += 1

        self.rows_written += n
        log.info(
            f"  [Requisiti] '{func_id}' (riga {func_row}): "
            f"{len(nuovi)} nuovi, {len(aggiornati)} aggiornati"
        )
        return n

    # ─────────────────────────────────────────────────────────────────────
    # Composizione delle voci
    # ─────────────────────────────────────────────────────────────────────

    def _merge(self, reqs_by_config: dict) -> list[dict]:
        """
        Unisce i requisiti trovati nelle varie configurazioni.

        REQ_ONE_ROW_PER_CONFIG = False → una riga per ID, col D con
            l'elenco colorato di tutte le configurazioni in cui compare.
        REQ_ONE_ROW_PER_CONFIG = True  → una riga per (ID, configurazione).
        """
        if cfg.REQ_ONE_ROW_PER_CONFIG:
            out = []
            for cname in sorted(reqs_by_config):
                for r in reqs_by_config[cname]:
                    out.append({"req_id": r.req_id, "configs": [cname], "req": r})
            return out

        merged: dict[str, dict] = {}
        for cname in sorted(reqs_by_config):
            for r in reqs_by_config[cname]:
                e = merged.setdefault(
                    _norm(r.req_id),
                    {"req_id": r.req_id, "configs": [], "req": r},
                )
                if cname not in e["configs"]:
                    e["configs"].append(cname)
                if _richness(r) > _richness(e["req"]):
                    e["req"] = r
        return list(merged.values())

    # ─────────────────────────────────────────────────────────────────────
    # Scrittura della singola riga di secondo livello
    # ─────────────────────────────────────────────────────────────────────

    def _fill_row(self, row: int, entry: dict, progressive: int | None) -> None:
        req_id  = entry["req_id"]
        configs = entry["configs"]
        primary = config_colors.color_for(configs[0]) if configs else "000000"

        # Col A e B restano VUOTE: è il discriminatore di secondo livello
        self.ws.cell(row, cfg.REQ_FUNC_ID_COL).value = None
        self.ws.cell(row, cfg.REQ_DESC_COL).value    = None

        # Col C — Nr requisito
        c = self.ws.cell(row, cfg.REQ_NUMBER_COL)
        c.value = f"{progressive}) {req_id}" if progressive else req_id
        c.font = Font(color=primary, bold=False)
        c.alignment = Alignment(wrap_text=True, vertical="top")

        # Col D — Configurazione di appartenenza, colorata
        d = self.ws.cell(row, cfg.REQ_CONFIG_COL)
        d.alignment = Alignment(wrap_text=True, vertical="top")
        if _RICH and len(configs) > 1:
            blocks = []
            for i, cname in enumerate(configs):
                if i:
                    blocks.append(TextBlock(InlineFont(color="000000"), ", "))
                blocks.append(
                    TextBlock(InlineFont(color=config_colors.color_for(cname)), cname)
                )
            d.value = CellRichText(*blocks)
        else:
            d.value = ", ".join(configs)
            d.font = Font(color=primary, bold=False)

        # Colonne extra opzionali (E, F, G...)
        if cfg.REQ_WRITE_EXTRA_FIELDS:
            req = entry["req"]
            col = cfg.REQ_EXTRA_FIELDS_START_COL
            for attr in ("description", "req_type", "derived_to",
                         "user_interface", "sil"):
                val = getattr(req, attr, "")
                if val:
                    cc = self.ws.cell(row, col)
                    cc.value = val
                    cc.alignment = Alignment(wrap_text=True, vertical="top")
                col += 1


# ---------------------------------------------------------------------------
# Helper
# ---------------------------------------------------------------------------

def _s(v) -> str:
    if v is None:
        return ""
    t = type(v).__name__
    if "Formula" in t or "formula" in t:
        return ""
    return str(v).strip()


def _norm(s: str) -> str:
    return " ".join((s or "").split()).lower()


def _strip_progressive(s: str) -> str:
    """Rimuove un eventuale prefisso progressivo '12) ' dal Nr requisito."""
    import re
    return re.sub(r"^\s*\d+\)\s*", "", s or "")


def _richness(r) -> int:
    return sum(bool(getattr(r, a, "")) for a in
               ("description", "req_type", "derived_to", "user_interface", "sil"))